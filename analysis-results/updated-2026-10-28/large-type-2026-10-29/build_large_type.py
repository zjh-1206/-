#!/usr/bin/env python3
"""Final-size typography: separate projection figures and compact paper plates."""
import argparse, hashlib, json, os, shutil, tempfile
from pathlib import Path
os.environ.setdefault('MPLCONFIGDIR',str(Path(tempfile.gettempdir())/'large_type_mpl'))
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.colors import Normalize
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from build_readable_graphs import parse,matrix,improve_project

NAVY='#15384c'; GREY='#586671'

def main():
    ap=argparse.ArgumentParser();ap.add_argument('project',type=Path);ap.add_argument('--output',required=True,type=Path);a=ap.parse_args()
    out=a.output;out.mkdir(parents=True,exist_ok=True);fd=out/'figures';fd.mkdir(exist_ok=True)
    raw=a.project.read_bytes();data,models=parse(raw)
    (out/'source_project.dxpx').write_bytes(raw)
    (out/'response_surface_readable.dxpx').write_bytes(improve_project(raw)[0])
    data.to_csv(out/'source_data.csv',index=False,float_format='%.15g')
    for name in ['build_large_type.py','build_readable_graphs.py']:
        src=Path(__file__).parent/name
        if src.resolve()!=(out/name).resolve():shutil.copy2(src,out/name)
    plt.rcParams.update({'font.family':'Liberation Serif','font.size':24,'axes.labelsize':28,
        'axes.titlesize':28,'xtick.labelsize':24,'ytick.labelsize':24,'legend.fontsize':23,
        'axes.unicode_minus':False,'axes.spines.top':False,'axes.spines.right':False,
        'axes.linewidth':1.3,'svg.fonttype':'none','pdf.fonttype':42,'savefig.dpi':240})
    v=np.linspace(-1,1,121);aa,bb=np.meshgrid(v,v);tt=85+10*aa;rr=1+.333*bb
    def predict(col,x):
        terms,coef,_=models[col];return matrix(x,terms)@coef
    def field(col,c):
        xx=np.column_stack([aa.ravel(),bb.ravel(),np.full(aa.size,c)])
        z=predict(col,xx).reshape(aa.shape)
        return np.ma.array(z,mask=np.abs(aa)+np.abs(bb)+abs(c)>2+1e-10)
    pr=Presentation();pr.slide_width=Inches(13.333);pr.slide_height=Inches(7.5)
    pdf=PdfPages(out/'large_type_figures.pdf');manifest=[];qa=[]
    def text(sl,x,y,w,h,t,size,colour=NAVY,bold=False):
        tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True
        tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=Inches(.01)
        p=tf.paragraphs[0];p.text=t
        for run in p.runs:
            run.font.name='Microsoft YaHei';run.font.size=Pt(size);run.font.bold=bold
            run.font.color.rgb=RGBColor.from_string(colour.lstrip('#'))
        return tb
    def validate_bounds(fig,stem):
        fig.canvas.draw();renderer=fig.canvas.get_renderer();bad=[]
        # All rendered text must fit the figure. The PPT title/subtitle are separate editable text.
        from matplotlib.text import Text
        for obj in fig.findobj(match=Text):
            if not obj.get_visible() or not obj.get_text().strip():continue
            box=obj.get_window_extent(renderer)
            # Matplotlib also creates hidden tick-label candidates outside explicit limits.
            if obj.get_clip_on():continue
            if box.x0 < -2 or box.y0 < -2 or box.x1 > fig.bbox.width+2 or box.y1 > fig.bbox.height+2:
                bad.append(obj.get_text())
        if bad:raise ValueError((stem,'out-of-figure labels',bad))
    def finish(fig,stem,title,subtitle,note,min_font=23,side=None):
        validate_bounds(fig,stem)
        png=fd/(stem+'.png');fig.savefig(png,facecolor='white')
        svg=fd/(stem+'.svg');fig.savefig(svg,facecolor='white')
        svg.write_text(svg.read_text().replace("'Liberation Serif'","'Times New Roman', 'Liberation Serif', serif"))
        pdf.savefig(fig,facecolor='white')
        wi,hi=fig.get_size_inches();plt.close(fig)
        sl=pr.slides.add_slide(pr.slide_layouts[6])
        text(sl,.45,.18,12.4,.6,title,32,bold=True)
        text(sl,.45,.91,12.4,.47,subtitle,22,colour=GREY)
        if side:
            x=.15;y=1.50;scale=1.0
            sl.shapes.add_picture(str(png),Inches(x),Inches(y),width=Inches(wi),height=Inches(hi))
            text(sl,8.88,2.03,3.95,3.9,side,24)
        else:
            maxw,maxh=12.45,5.65
            scale=min(maxw/wi,maxh/hi)
            pw,ph=wi*scale,hi*scale
            sl.shapes.add_picture(str(png),Inches((13.333-pw)/2),Inches(1.46+(5.65-ph)/2),width=Inches(pw),height=Inches(ph))
        final_size=min_font*scale
        assert final_size>=22-1e-8,(stem,final_size)
        sl.notes_slide.notes_text_frame.text=(note+'\n来源：用户上传DX13原项目，F二次/P 2FI。'
          '\nPNG/PDF为Liberation Serif替代字体；SVG请求Times New Roman。'
          f'\n图内主要数字/图例最小最终字号{final_size:.2f}pt。')
        manifest.append({'file':stem,'slide':len(pr.slides),'title':title,'subtitle':subtitle})
        qa.append({'figure':stem,'source_width_in':float(wi),'source_height_in':float(hi),
                   'placement_scale':float(scale),'minimum_source_font_pt':min_font,
                   'minimum_final_font_pt':float(final_size),'axis_title_source_font_pt':26 if side else 28})
    def style(ax):
        ax.set(xlim=(75,95),ylim=(.667,1.333),xlabel='Temperature (°C)',ylabel='Phase ratio')
        ax.set_xticks([75,80,85,90,95]);ax.set_yticks([.667,1,1.333],['0.667','1.000','1.333'])
        ax.tick_params(length=6,pad=7);ax.set_facecolor('#eceff1')
        ax.xaxis.labelpad=11;ax.yaxis.labelpad=15

    # Six individually readable contours replace the small three-panel pages.
    for col,name,scale,levels in [('F','氟萃取率',(75,90),np.arange(76,90,2)),
                                 ('P','磷萃取率',(0,11),np.arange(2,11,2))]:
        for c,code in [(-1.,'0p8'),(0.,'1p0'),(1.,'1p2')]:
            fig=plt.figure(figsize=(12.2,5.6));ax=fig.add_axes([.14,.21,.66,.75])
            z=field(col,c)
            cf=ax.contourf(tt,rr,z,levels=np.linspace(*scale,61),cmap='viridis')
            cs=ax.contour(tt,rr,z,levels=levels,colors='#173442',linewidths=1.2)
            labels=ax.clabel(cs,fmt='%g',fontsize=24,inline=True,inline_spacing=7)
            for label in labels:label.set_bbox({'facecolor':'white','edgecolor':'none','alpha':.8,'pad':1.0})
            subset=data[np.isclose(data.Initial_F,1+.2*c)]
            ax.scatter(subset.Temperature,subset.Ratio,s=65,facecolor='white',edgecolor=NAVY,
                       linewidth=1.2,zorder=5,clip_on=False)
            style(ax)
            cax=fig.add_axes([.84,.21,.024,.75])
            cb=fig.colorbar(cf,cax=cax,ticks=[75,80,85,90] if col=='F' else [0,2,4,6,8,10])
            cb.set_label(('F' if col=='F' else 'P')+' extraction (%)',fontsize=28,labelpad=12)
            cb.ax.tick_params(labelsize=24,pad=6)
            sub=f'初始氟含量 = {1+.2*c:.1f}%　｜　统一色标；圆点为已测工况位置'
            if c:sub=f'初始氟含量 = {1+.2*c:.1f}%　｜　灰区超出实验工况凸包'
            finish(fig,f'{col}_contour_C{code}',name+'随温度和相比的变化',sub,
                   '色面表示保存模型条件均值预测。相同响应各页采用相同坐标、色标、等高线。灰色区域不作同等实验支持的预测。',24)

    # Separate each response, simplify legend, and use one temperature chart per page.
    colors=['#2166ac','#087f8c','#c86625']
    for col,name in [('F','氟萃取率'),('P','磷萃取率')]:
        fig=plt.figure(figsize=(12.2,5.6));ax=fig.add_axes([.14,.21,.82,.75])
        for c,color in zip([-1.,0.,1.],colors):
            x=np.column_stack([v,np.zeros(len(v)),np.full(len(v),c)])
            ax.plot(85+10*v,predict(col,x),color=color,lw=3.2,label=f'{1+.2*c:.1f}%')
        ax.set(xlabel='Temperature (°C)',ylabel=('F' if col=='F' else 'P')+' extraction (%)',
               xlim=(75,95),ylim=(80,90) if col=='F' else (4,7.1))
        ax.set_xticks([75,80,85,90,95]);ax.tick_params(length=6,pad=7)
        ax.xaxis.labelpad=11;ax.yaxis.labelpad=15;ax.grid(alpha=.17)
        ax.legend(title='Initial fluoride',fontsize=24,title_fontsize=24,frameon=False,
                  loc='upper left' if col=='F' else 'upper right',ncols=3)
        finish(fig,col+'_temperature_trends','初始氟含量如何影响'+name+'的温度响应',
               '相比固定为 1.000　｜　曲线为模型条件均值预测',
               'F使用二次模型、P保持原项目2FI。斜率不直接等于显著性；各响应纵轴不同，不横跨响应比较曲线视觉斜率。',24)

    # Single large tradeoff graph, shortened and enlarged annotations.
    bv=np.linspace(-1,1,401);xx=np.column_stack([np.zeros(len(bv)),bv,np.zeros(len(bv))])
    ff,pp=predict('F',xx),predict('P',xx)
    fig=plt.figure(figsize=(12.2,5.6));ax=fig.add_axes([.13,.21,.67,.75])
    ax.fill_between([0,5],80,91,color='#e5f2e8');ax.axhline(80,color=GREY,ls='--',lw=1.5);ax.axvline(5,color=GREY,ls='--',lw=1.5)
    ax.plot(pp,ff,color=GREY,lw=1.8);sc=ax.scatter(pp,ff,c=1+.333*bv,cmap='viridis',s=18,zorder=2)
    for ratio,offset in [(.667,(9,9)),(.830,(9,-29)),(1.,(9,-30)),(1.333,(-89,-28))]:
        x=np.array([[0,(ratio-1)/.333,0]])
        f,p=float(predict('F',x)[0]),float(predict('P',x)[0])
        ax.scatter(p,f,s=75,facecolor='white',edgecolor=NAVY,zorder=4)
        ax.annotate(f'{ratio:.3f}',(p,f),xytext=offset,textcoords='offset points',fontsize=24)
    ax.text(.3,90.1,'F ≥ 80%, P ≤ 5%',fontsize=24,color='#386b53')
    ax.set(xlabel='P extraction (%)',ylabel='F extraction (%)',xlim=(0,11),ylim=(76,91))
    ax.set_xticks([0,2,4,6,8,10]);ax.set_yticks([76,80,84,88]);ax.tick_params(length=6,pad=7)
    ax.xaxis.labelpad=11;ax.yaxis.labelpad=15;ax.grid(alpha=.13)
    cb=fig.colorbar(sc,cax=fig.add_axes([.84,.21,.024,.75]),ticks=[.667,1,1.333])
    cb.ax.set_yticklabels(['0.667','1.000','1.333']);cb.set_label('Phase ratio',fontsize=28,labelpad=12)
    finish(fig,'FP_phase_ratio_tradeoff','提高相比时，氟回收与磷损失同步变化',
           '温度 = 85°C；初始氟含量 = 1.0%　｜　点旁数字为相比',
           '绿色区域满足F≥80、P≤5模型均值目标，不能保证未来每一批均达标。原始模型与实验数据未改变。',24)

    # Optional 3D diagrams stay separate, with readable side explanations.
    for c,code in [(0.,'1p0'),(-1.,'0p8')]:
        z=field('F',c);fig=plt.figure(figsize=(8.5,5.6));ax=fig.add_subplot(111,projection='3d')
        fig.subplots_adjust(left=.04,right=.95,bottom=.12,top=.98)
        ax.plot_surface(tt,rr,z.filled(np.nan),cmap='viridis',norm=Normalize(75,90),
                        rcount=45,ccount=45,linewidth=.15,edgecolor='#53656a',alpha=.98)
        ax.set(xlim=(75,95),ylim=(.667,1.333),zlim=(75,90))
        ax.set_xticks([75,85,95]);ax.set_yticks([1,1.333]);ax.set_zticks([75,80,85,90])
        ax.tick_params(axis='x',labelsize=23,pad=0);ax.tick_params(axis='y',labelsize=23,pad=8);ax.tick_params(axis='z',labelsize=23,pad=5)
        ax.view_init(26,-128);ax.set_box_aspect((1.55,1.05,.85))
        # Short axis symbols avoid intersections between large 3D titles and ticks.
        ax.set_xlabel('A',fontsize=26,labelpad=18)
        ax.set_ylabel('B',fontsize=26,labelpad=30)
        ax.set_zlabel('F',fontsize=26,labelpad=16)
        side='轴标识\n\nA：温度 / °C\nB：相比\nF：氟萃取率 / %\n\nF 轴：75–90%\n曲面为模型预测。'
        if c:side+='\n仅显示设计凸包内区域。'
        finish(fig,'F_surface_C'+code,'氟萃取率曲面：观察整体变化方向',
               f'初始氟含量 = {1+.2*c:.1f}%　｜　各曲面采用相同范围和视角',
               'Z轴截取75–90%用于显示该实验范围内差异；视觉坡度不能直接表示绝对效应。',23,side)
    pdf.close();pr.save(out/'large_type_slides.pptx')

    # A distinct 180 mm paper-layout plate. These small print fonts are not used in the PPT.
    with plt.rc_context({'font.size':10,'axes.labelsize':11,'axes.titlesize':11,
                         'xtick.labelsize':10,'ytick.labelsize':10,'axes.linewidth':.8}):
        fig,axs=plt.subplots(2,3,figsize=(180/25.4,132/25.4))
        fig.subplots_adjust(left=.1,right=.87,bottom=.12,top=.92,wspace=.32,hspace=.52)
        for row,col in enumerate(['F','P']):
            scale=(75,90) if col=='F' else (0,11)
            for j,c in enumerate([-1.,0.,1.]):
                ax=axs[row,j];z=field(col,c)
                filled=ax.contourf(tt,rr,z,levels=np.linspace(*scale,61),cmap='viridis')
                cs=ax.contour(tt,rr,z,levels=np.arange(78,90,4) if col=='F' else [2,5,8],colors=NAVY,linewidths=.55)
                ax.clabel(cs,fmt='%g',fontsize=10)
                ax.set(xlim=(75,95),ylim=(.667,1.333),xlabel='Temperature (°C)',ylabel='Phase ratio' if j==0 else '',
                       title=f'({chr(97+row*3+j)}) C = {1+.2*c:.1f}%')
                ax.set_xticks([75,85,95]);ax.set_yticks([.667,1,1.333],['0.667','1','1.333']);ax.set_facecolor('#eceff1')
            cax=fig.add_axes([.9,.565 if row==0 else .12,.018,.355])
            cb=fig.colorbar(filled,cax=cax,ticks=[75,80,85,90] if col=='F' else [0,5,10]);cb.set_label(col+' extraction (%)',fontsize=11)
        fig.savefig(out/'paper_layout_180mm.pdf',facecolor='white')
        fig.savefig(out/'paper_layout_180mm.png',dpi=400,facecolor='white')
        plt.close(fig)
    pd.DataFrame(manifest).to_csv(out/'figure_index.csv',index=False,encoding='utf-8-sig')
    (out/'font_size_audit.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
    refs='''# 格式参考与字号依据

## 实际可用的视觉参考
用户此前上传的论文图“图11 不同光照条件下净光合速率的等高线图”。仅依据可见图片：统一坐标与色标、固定第三因素、分面、清楚标注单位；文章题名、作者、DOI未知，未编造书目信息。

## 本轮官方资料访问记录
Nature、PLOS ONE、Elsevier、MDPI及PMC在本环境均返回网络代理403，本轮没有成功读取正文。因此本包不声称其推荐了任何具体字号，也不声称已查阅若干新论文。
待核实的官方来源：https://journals.plos.org/plosone/s/figures ；https://www.nature.com/nature/for-authors/final-submission 。

## 为什么同时提供两种尺寸
PPT采用实际放置尺寸，坐标标题26–28pt，刻度、等高线、图例23–24pt，标题32pt；一页一张主图，长解释放备注。所有主要数字最终字号>=22pt，见font_size_audit.json。
论文排版示例采用180mm宽、坐标标题11pt、刻度与等高线10pt。此为本次设计选择，不是核实过的某期刊强制规定。论文版需按具体投稿期刊规范再次调整，不应用其小字号替代PPT大字图。

## 字体
当前环境没有真正的Times New Roman，PNG/PDF使用Liberation Serif作为替代；SVG保留可编辑文字并请求Times New Roman。不能据此声称PNG/PDF已经是新罗马。PPT中文标题请求Microsoft YaHei。
'''
    (out/'format_references.md').write_text(refs)
    readme='''# 大字版响应面图（2026-10-29）

本版针对“字体太小”重新排版，不是只增大源文件中的字号。原三联图拆为单图，F/P温度趋势分别一页；主要刻度、等高线和图例最终字号至少22pt，坐标名26–28pt。

- large_type_slides.pptx：11页图版，可按汇报需要选用。
- large_type_figures.pdf：11页图形，最终字号与PPT对应。
- figures：高清PNG及可编辑文字SVG。
- paper_layout_180mm.pdf/png：单独的180mm宽论文排版示例，字小于PPT版，供论文排版使用。
- font_size_audit.json：每张图最终字号和缩放倍数。
- format_references.md：格式参考、实际访问限制和字体说明。
- response_surface_readable.dxpx：与上一版相同的原生显示调整副本，字体偏好保存在用户软件内，尚未在DX13中试开。

图形全部使用上传项目完整精度数据及保存的F二次/P 2FI模型，数据、模型与优化设置不变。色面/曲线为条件均值预测；各固定C水平使用统一坐标和色标，灰区超出实验工况凸包。β数值版本差异仍按原项目保留，不据其作可靠优化。没有实际能耗数据。

复现：安装requirements.txt依赖后，运行 `python build_large_type.py source_project.dxpx --output regenerated`。
'''
    (out/'README.md').write_text(readme)
    import importlib.metadata
    (out/'requirements.txt').write_text('\n'.join(p+'=='+importlib.metadata.version(p) for p in ['numpy','pandas','scipy','statsmodels','matplotlib','python-pptx','Pillow'])+'\n')
    print(json.dumps({'figures':len(manifest),'minimum_final_font':min(q['minimum_final_font_pt'] for q in qa),'output':str(out)}))

if __name__=='__main__':main()
