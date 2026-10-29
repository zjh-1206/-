#!/usr/bin/env python3
"""Assemble the reviewed figure set into six group-meeting slides."""
from pathlib import Path
import argparse
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image

def main():
    ap=argparse.ArgumentParser();ap.add_argument('folder',type=Path);a=ap.parse_args()
    out=a.folder;pr=Presentation();pr.slide_width=Inches(13.333);pr.slide_height=Inches(7.5)
    items=[
      ('温度与相比如何共同影响氟萃取率','01_F_surface_C1p0',
       '初始氟含量固定1.0%；曲面和等高线对应同一模型。Z轴75–90%，用于展示实验范围内差异。'),
      ('初始氟含量0.8%：恢复完整因素范围','01_F_surface_C0p8',
       '灰色角落超出实验设计凸包；这部分不作有同等实验支持的预测展示。'),
      ('比较三个初始氟水平下的氟萃取趋势','02_F_three_conditions',
       '三个面板采用相同坐标、色标和等高线。圆点表示已测工况位置，色面表示模型预测。'),
      ('相比提高时，也需要关注磷损失','02_P_three_conditions',
       '使用上传项目保存的2FI模型；当前呈现与原项目对应，没有替换为先前报告的线性模型。'),
      ('温度的作用取决于初始氟含量','03_temperature_interaction',
       '相比固定1；曲线为模型条件均值。两面板纵轴尺度不同，斜率本身不表示统计显著性。'),
      ('用同一张图展示氟回收与磷损失的权衡','04_FP_phase_ratio_tradeoff',
       '固定85°C、初始氟1.0%。绿色区域满足模型均值F≥80%、P≤5%，仍需独立重复验证。'),
    ]
    navy=RGBColor.from_string('15384C');grey=RGBColor.from_string('586671')
    def text(sl,x,y,w,h,t,size=20,colour=navy):
        box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=box.text_frame;tf.word_wrap=True
        p=tf.paragraphs[0];p.text=t
        for r in p.runs:r.font.name='Microsoft YaHei';r.font.size=Pt(size);r.font.color.rgb=colour
        return box
    for n,(title,stem,note) in enumerate(items,1):
        sl=pr.slides.add_slide(pr.slide_layouts[6])
        text(sl,.48,.22,12.3,.65,title,28)
        p=out/'figures'/(stem+'.png');im=Image.open(p)
        width,height=12.8,5.6
        scale=min(width/im.width,height/im.height)
        wi,hi=im.width*scale,im.height*scale
        sl.shapes.add_picture(str(p),Inches((13.333-wi)/2),Inches(1.0+(5.6-hi)/2),width=Inches(wi),height=Inches(hi))
        text(sl,.5,6.72,12.3,.52,note,16,grey)
        sl.notes_slide.notes_text_frame.text=(note+'\n图源：用户上传的响应面设计-1.dxpx。原数据、模型和优化设置保持。'
            '\nPNG图内使用Liberation Serif替代预览；可编辑SVG请求Times New Roman。项目字体继承软件全局偏好。')
    pr.save(out/'readable_graphs.pptx')

if __name__=='__main__':main()
